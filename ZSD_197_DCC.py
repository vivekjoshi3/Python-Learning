import pandas as pd
from openpyxl import load_workbook

# --------------------------------------------------------------------
# FILE PATHS  (EDIT THESE)
# --------------------------------------------------------------------
report_file = "DCC Report HW.xlsx"          # Destination file
source_file = "ZSD_197_DCC.xlsx"            # Source file

source_sheet_name = "Data"                  # Sheet to read from source
target_sheet_name = "ZSD_197_DCC"           # Sheet to modify in report
# --------------------------------------------------------------------

# Read source file
df = pd.read_excel(source_file, sheet_name=source_sheet_name)

# Load report workbook
wb = load_workbook(report_file)
ws = wb[target_sheet_name]   # Only touching this sheet

# -----------------------------------------------------------
# 1️⃣ REMOVE FILTERS (if any)
# -----------------------------------------------------------
ws.auto_filter.ref = None    # Removes existing filter area if present

# -----------------------------------------------------------
# 2️⃣ UNHIDE ALL COLUMNS ONLY
# -----------------------------------------------------------
for col_letter in ws.column_dimensions:
    ws.column_dimensions[col_letter].hidden = False

# -----------------------------------------------------------
# 3️⃣ CLEAR COLUMNS I to BD (9 → 56)
# -----------------------------------------------------------
start_col = 9
end_col = 56

for col in range(start_col, end_col + 1):
    for row in range(1, ws.max_row + 1):
        ws.cell(row=row, column=col).value = None

# -----------------------------------------------------------
# 4️⃣ WRITE NEW HEADERS + DATA STARTING AT COLUMN I
# -----------------------------------------------------------

# Write headers (first row)
for c_idx, col_name in enumerate(df.columns, start=start_col):
    ws.cell(row=1, column=c_idx).value = col_name

# Write data (starting from row 2)
for r_idx, row in enumerate(df.itertuples(index=False), start=2):
    for c_idx, value in enumerate(row, start=start_col):
        ws.cell(row=r_idx, column=c_idx).value = value

# -----------------------------------------------------------
# Save workbook
# -----------------------------------------------------------
wb.save(report_file)

print("✔ Completed: Headers replaced, data updated, filters removed, columns unhidden.")
